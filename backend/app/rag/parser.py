import io
import logging


logger = logging.getLogger(__name__)


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text content from PDF file bytes using PyMuPDF (fitz)."""
    try:
        import fitz  # PyMuPDF
    except ImportError:
        raise RuntimeError("PyMuPDF is not installed in the environment.")

    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        text_parts = []
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text()
            if text:
                text_parts.append(text.strip())
        doc.close()
        full_text = "\n\n".join(text_parts).strip()
        if not full_text:
            raise ValueError("PDF file contains no readable text or is scanned without OCR.")
        return full_text
    except Exception as e:
        logger.error(f"Error parsing PDF file: {e}")
        if isinstance(e, ValueError):
            raise e
        raise ValueError(f"Failed to parse PDF file: {str(e)}")


def extract_text_from_pptx(file_bytes: bytes) -> str:
    """Extract text content from PPTX file bytes using python-pptx."""
    try:
        from pptx import Presentation
    except ImportError:
        raise RuntimeError("python-pptx is not installed in the environment.")

    try:
        prs = Presentation(io.BytesIO(file_bytes))
        text_parts = []
        for slide_idx, slide in enumerate(prs.slides, start=1):
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    clean = shape.text.strip()
                    if clean:
                        slide_text.append(clean)
            if slide_text:
                text_parts.append(f"Slide {slide_idx}:\n" + "\n".join(slide_text))

        full_text = "\n\n".join(text_parts).strip()
        if not full_text:
            raise ValueError("PPTX file contains no readable slide text.")
        return full_text
    except Exception as e:
        logger.error(f"Error parsing PPTX file: {e}")
        if isinstance(e, ValueError):
            raise e
        raise ValueError(f"Failed to parse PPTX presentation: {str(e)}")


def extract_text_from_txt(file_bytes: bytes) -> str:
    """Extract text content from TXT file bytes."""
    for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
        try:
            text = file_bytes.decode(encoding).strip()
            if text:
                return text
        except Exception:
            continue
    raise ValueError("TXT file is empty or could not be decoded with standard text encodings.")


def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    """
    Dispatcher to extract clean text based on file extension.
    Supports PDF, PPTX, TXT.
    """
    if not file_bytes or len(file_bytes) == 0:
        raise ValueError("Uploaded file is empty (0 bytes).")

    ext = filename.lower().split(".")[-1] if "." in filename else ""
    if ext == "pdf":
        return extract_text_from_pdf(file_bytes)
    elif ext == "pptx":
        return extract_text_from_pptx(file_bytes)
    elif ext in ["txt", "text", "md", "csv"]:
        return extract_text_from_txt(file_bytes)
    else:
        raise ValueError(f"Unsupported file format '.{ext}'. Supported formats are: PDF, PPTX, TXT.")
